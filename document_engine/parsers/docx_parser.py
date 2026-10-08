import io
import logging
from typing import List, Dict, Any, Tuple
from docx import Document
from .base import BaseParser
from ..schemas.document import PageModel, TableModel

logger = logging.getLogger(__name__)

class DOCXParser(BaseParser):
    """Parses Microsoft Word (.docx) documents preserving headings, paragraphs, and tables."""

    def parse(self, file_bytes: bytes, filename: str, document_id: str) -> Tuple[List[PageModel], Dict[str, Any]]:
        doc = Document(io.BytesIO(file_bytes))
        
        # Extract tables
        tables: List[TableModel] = []
        for t_idx, table in enumerate(doc.tables):
            headers = []
            rows = []
            if len(table.rows) > 0:
                headers = [cell.text.strip() for cell in table.rows[0].cells]
                for r in table.rows[1:]:
                    rows.append([cell.text.strip() for cell in r.cells])
            
            # Format markdown table
            md_lines = []
            if headers:
                md_lines.append("| " + " | ".join(headers) + " |")
                md_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                for r in rows:
                    md_lines.append("| " + " | ".join(r) + " |")
            table_md = "\n".join(md_lines)

            tables.append(TableModel(
                page_number=1,
                headers=headers,
                rows=rows,
                markdown=table_md,
                metadata={"table_index": t_idx + 1}
            ))

        # Extract text blocks
        paragraphs_text = []
        headings = []
        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt:
                if p.style.name.startswith("Heading"):
                    headings.append(txt)
                    paragraphs_text.append(f"\n### {txt}\n")
                else:
                    paragraphs_text.append(txt)

        # Include table markdown in full page text
        full_text_blocks = paragraphs_text.copy()
        if tables:
            full_text_blocks.append("\n\n### Document Tables:\n")
            for t in tables:
                if t.markdown:
                    full_text_blocks.append(t.markdown + "\n")

        full_text = "\n\n".join(full_text_blocks).strip()

        # Build PageModel (Word documents default to single or section-based pages)
        page = PageModel(
            document_id=document_id,
            page_number=1,
            text=full_text,
            tables=tables,
            metadata={
                "headings": headings,
                "paragraph_count": len(doc.paragraphs),
                "table_count": len(doc.tables),
                "char_count": len(full_text)
            }
        )

        doc_metadata = {
            "parser": "DOCXParser",
            "page_count": 1,
            "headings_count": len(headings),
            "tables_count": len(tables)
        }

        return [page], doc_metadata
