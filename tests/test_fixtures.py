import io
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
import fitz # PyMuPDF
import docx
import openpyxl
import pandas as pd

def create_sample_pdf(filepath: Path, content: str = None) -> bytes:
    """Creates a sample digital PDF document."""
    doc = fitz.open()
    page = doc.new_page()
    text = content or (
        "LOAN APPLICATION FORM - APEX FINANCIAL SERVICES\n\n"
        "Applicant Name: Ramesh Kumar\n"
        "Date of Birth: 1985-06-15\n"
        "PAN Number: ABCDE1234F\n"
        "Aadhaar Number: 9876 5432 1098\n"
        "Employer: Acme Global Technologies Ltd\n"
        "Monthly Income: 45000 INR\n"
        "Loan Amount: 500000 INR\n"
        "Status: Under Review\n"
        "Contact: +91 9876543210\n"
        "Email: ramesh.kumar@example.com\n"
    )
    page.insert_text((50, 72), text, fontsize=12)
    pdf_bytes = doc.tobytes()
    if filepath:
        filepath.write_bytes(pdf_bytes)
    return pdf_bytes

def create_scanned_pdf(filepath: Path) -> bytes:
    """Creates a scanned PDF where pages contain only images with no direct digital text."""
    doc = fitz.open()
    page = doc.new_page()
    
    # Create an image containing text
    img = Image.new("RGB", (600, 800), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((50, 50), "SALARY CERTIFICATE (SCANNED STAMPED COPY)", fill=(0, 0, 0))
    draw.text((50, 100), "Employee: Ramesh Kumar", fill=(0, 0, 0))
    draw.text((50, 150), "Organization: Acme Global Technologies Ltd", fill=(0, 0, 0))
    draw.text((50, 200), "Monthly Income: 45000 INR", fill=(0, 0, 0))
    
    img_byte_arr = io.BytesIO()
    img.save(img_byte_arr, format="PNG")
    page.insert_image(page.rect, stream=img_byte_arr.getvalue())
    
    pdf_bytes = doc.tobytes()
    if filepath:
        filepath.write_bytes(pdf_bytes)
    return pdf_bytes

def create_sample_docx(filepath: Path) -> bytes:
    """Creates a sample Word document with headings, paragraphs, and tables."""
    doc = docx.Document()
    doc.add_heading("EMPLOYMENT VERIFICATION LETTER", 0)
    doc.add_paragraph("This is to certify that Mr. Ramesh Kumar is employed with Acme Global Technologies Ltd.")
    
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Designation"
    hdr_cells[1].text = "Monthly Salary"
    hdr_cells[2].text = "Joining Date"
    
    row_cells = table.add_row().cells
    row_cells[0].text = "Senior Software Engineer"
    row_cells[1].text = "45000 INR"
    row_cells[2].text = "2021-04-01"
    
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()
    if filepath:
        filepath.write_bytes(docx_bytes)
    return docx_bytes

def create_sample_xlsx(filepath: Path) -> bytes:
    """Creates a sample Excel financial statement."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Payroll & Disbursals"
    
    ws.append(["Month", "Applicant", "Income", "Tax Deducted", "Net Disbursed"])
    ws.append(["Jan 2026", "Ramesh Kumar", 45000, 3000, 42000])
    ws.append(["Feb 2026", "Ramesh Kumar", 45000, 3000, 42000])
    ws.append(["Mar 2026", "Ramesh Kumar", 45000, 3000, 42000])
    
    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()
    if filepath:
        filepath.write_bytes(xlsx_bytes)
    return xlsx_bytes

def create_sample_csv(filepath: Path) -> bytes:
    """Creates a sample CSV bank statement."""
    csv_content = (
        "Transaction Date,Description,Credit,Debit,Balance,Entity\n"
        "2026-01-31,Salary from Acme Global,42000,0,125000,Ramesh Kumar\n"
        "2026-02-28,Salary from Acme Global,42000,0,167000,Ramesh Kumar\n"
        "2026-03-31,Salary from Acme Global,42000,0,209000,Ramesh Kumar\n"
    )
    csv_bytes = csv_content.encode("utf-8")
    if filepath:
        filepath.write_bytes(csv_bytes)
    return csv_bytes

def create_sample_image(filepath: Path) -> bytes:
    """Creates a standalone sample image with ID card text."""
    img = Image.new("RGB", (500, 300), color=(240, 240, 240))
    draw = ImageDraw.Draw(img)
    draw.text((20, 20), "GOVERNMENT IDENTITY CARD", fill=(0, 0, 0))
    draw.text((20, 60), "Name: Ramesh Kumar", fill=(0, 0, 0))
    draw.text((20, 100), "PAN: ABCDE1234F", fill=(0, 0, 0))
    draw.text((20, 140), "DOB: 1985-06-15", fill=(0, 0, 0))
    
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    img_bytes = buf.getvalue()
    if filepath:
        filepath.write_bytes(img_bytes)
    return img_bytes
